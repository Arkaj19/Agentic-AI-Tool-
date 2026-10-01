"""
Row-by-row validation Excel (ported from sap-migration-agent tools/excel_report.py).

Reads the row-validation table from Databricks through the SQL warehouse and writes:

    Summary      overall status, row counts, failing rows per rule, rule results
    Failed Rows  only the rows that failed (red), failing cells in strong red
    All Rows     every row; failed rows light red, failing cells strong red
"""

import os
from collections import Counter
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

from app.databricks import client

RED_ROW = PatternFill("solid", fgColor="FFC7CE")      # light red: failed row
RED_CELL = PatternFill("solid", fgColor="FF6B6B")     # strong red: failing cell
GREEN = PatternFill("solid", fgColor="C6EFCE")
HEADER = PatternFill("solid", fgColor="1F4E78")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
DARK_RED = Font(color="9C0006", bold=True)
DARK_GREEN = Font(color="006100", bold=True)


def _fetch(table: str, max_rows: int) -> tuple[list[str], list[list]]:
    return client.query(f"SELECT * FROM {table} ORDER BY _row_no LIMIT {int(max_rows)}")


def _header(ws, names):
    ws.append([_styled(ws, n, fill=HEADER, font=WHITE_BOLD) for n in names])


def _styled(ws, value, fill=None, font=None):
    cell = WriteOnlyCell(ws, value=value)
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    return cell


def _data_sheet(wb, name, columns, rows, idx):
    ws = wb.create_sheet(name)
    ws.freeze_panes = "E2"                       # keep row no/status/rules/reasons + header visible
    for i, col in enumerate(columns, 1):
        width = 55 if col == "_failure_reasons" else min(max(len(col) + 2, 10), 30)
        ws.column_dimensions[_col_letter(i)].width = width
    _header(ws, columns)
    for row in rows:
        failed = row[idx["_row_status"]] == "FAIL"
        bad_cols = set((row[idx["_failed_columns"]] or "").split(",")) if failed else set()
        cells = []
        for col, value in zip(columns, row):
            if col == "_row_status":
                cells.append(_styled(ws, value, RED_CELL if failed else GREEN,
                                     DARK_RED if failed else DARK_GREEN))
            elif col == "_failure_reasons" and failed:
                c = _styled(ws, value, RED_ROW)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                cells.append(c)
            elif failed and col in bad_cols:
                cells.append(_styled(ws, value, RED_CELL, DARK_RED))
            elif failed:
                cells.append(_styled(ws, value, RED_ROW))
            else:
                cells.append(value)
        ws.append(cells)
    ws.auto_filter.ref = f"A1:{_col_letter(len(columns))}{len(rows) + 1}"
    return ws


def _col_letter(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def export_row_validation_excel(row_table: str, out_path: str, rule_results: list[dict] = None,
                                title: str = "", max_rows: int = 200000) -> str:
    columns, rows = _fetch(row_table, max_rows)
    return write_excel(columns, rows, out_path, rule_results, title, row_table)


def write_excel(columns, rows, out_path, rule_results=None, title="", source=""):
    idx = {c: i for i, c in enumerate(columns)}
    for needed in ("_row_status", "_failed_rules", "_failed_columns"):
        if needed not in idx:
            raise ValueError(f"Row-validation table has no column {needed}")

    failed = [r for r in rows if r[idx["_row_status"]] == "FAIL"]
    per_rule = Counter(rid for r in failed for rid in (r[idx["_failed_rules"]] or "").split(";") if rid)
    status = "PASS" if not failed else "FAIL"

    wb = Workbook(write_only=True)

    # ---- Summary --------------------------------------------------------------------
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 16
    ws.append([_styled(ws, f"Row-by-row validation report {title}".strip(), font=Font(bold=True, size=14))])
    ws.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws.append(["Source table", source])
    ws.append(["Overall row status", _styled(ws, status, RED_CELL if failed else GREEN,
                                             DARK_RED if failed else DARK_GREEN)])
    ws.append(["Total rows", len(rows)])
    ws.append(["Passed rows", len(rows) - len(failed)])
    ws.append(["Failed rows", _styled(ws, len(failed), RED_ROW if failed else None)])
    ws.append([])
    ws.append([_styled(ws, "Legend", font=Font(bold=True))])
    ws.append([_styled(ws, "", RED_ROW), "Row failed at least one rule"])
    ws.append([_styled(ws, "", RED_CELL), "Cell that caused the failure"])
    ws.append([])

    if rule_results:
        keys = list(rule_results[0].keys())
        ws.append([_styled(ws, "Rule results (overall)", font=Font(bold=True))])
        _header(ws, keys + ["failing_rows (row check)"])
        for r in rule_results:
            rid = str(r.get("rule_id") or r.get("id") or "")
            st = str(r.get("status", "")).upper()
            vals = [r.get(k) if isinstance(r.get(k), (int, float, str, type(None))) else str(r.get(k))
                    for k in keys]
            vals = [_styled(ws, v, RED_CELL, DARK_RED) if k == "status" and st != "PASS"
                    else _styled(ws, v, GREEN, DARK_GREEN) if k == "status" else v
                    for k, v in zip(keys, vals)]
            ws.append(vals + [per_rule.get(rid, 0)])
    else:
        ws.append([_styled(ws, "Failing rows per rule", font=Font(bold=True))])
        _header(ws, ["rule_id", "failing_rows"])
        for rid, n in sorted(per_rule.items()):
            ws.append([rid, _styled(ws, n, RED_ROW)])
        if not per_rule:
            ws.append(["All rules", _styled(ws, "0 failing rows", GREEN, DARK_GREEN)])

    # ---- Data sheets -------------------------------------------------------------------
    _data_sheet(wb, "Failed Rows", columns, failed, idx)
    _data_sheet(wb, "All Rows", columns, rows, idx)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return out_path
